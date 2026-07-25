"""
Genera un Excel "desordenado" simulando cómo Lupita (dueña de taquería)
lleva sus registros actualmente: datos mezclados, sin estructura,
difícil de consultar. Usado como prop visual en la presentación.
"""
import random
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()

# ============================================================
# HOJA 1: "Ventas julio" — desordenada, con errores, notas al margen
# ============================================================
ws = wb.active
ws.title = "Ventas julio"

# Headers irregulares (como los pondría alguien sin formación)
headers = ["fecha", "que vendí", "cliente", "cuanto", "pago", "notas"]
ws.append(headers)

# Datos de taquería — desordenados, con inconsistencias intencionales
ventas_data = [
    ("22/07/2026", "5 tacos pastor + 2 aguas", "Don Pedro", 130, "efectivo", ""),
    ("22/07/2026", "3 quesadillas + refresco", "", 95, "efectivo", "no recuerdo el nombre"),
    ("21/07/2026", "orden familiar (12 tacos)", "Familia Lopez", 280, "transferencia", "pagaron al dia sgte"),
    ("23/07/2026", "2 tacos bistec + agua", "Señora del 7", 70, "efvo", ""),
    ("20/07/2026", "torta de milanesa + jugo", "Carlos", 85, "tarjeta", ""),
    ("22/07/2026", "8 tacos surtidos", "grupo oficina", 200, "transferencia", "pidieron factura??"),
    ("19/07/2026", "3 aguas + 2 refrescos", "", 65, "$", "solo bebidas"),
    ("23/07/2026", "5 tacos + 1 gringa + horchata", "Miguel", 155, "efectivo", ""),
    ("21/07/2026", "4 quesadillas + 3 tacos", "Ana", 175, "efectivo", "propina $20"),
    ("20/07/2026", "orden para llevar 10 tacos", "Sra Martinez", 230, "transf", ""),
    ("23/07/2026", "2 gringas + 2 aguas", "Roberto", 110, "efectivo", ""),
    ("19/07/2026", "6 tacos pastor", "????", 150, "efect", "no anote nombre"),
    ("22/07/2026", "torta + agua + tacos(3)", "Laura", 135, "tarjeta", "cobré de más? checar"),
    ("21/07/2026", "15 tacos evento", "Lic. Gomez", 350, "transferencia", "pago parcial $200"),
    ("23/07/2026", "4 tacos + horchata", "", 115, "efect", ""),
    ("20/07/2026", "2 ordenes quesadillas", "vecina Paty", 120, "efectivo", "le debo cambio $5"),
    ("18/07/2026", "comida empleados", "---", 0, "interno", "no cobrar"),
    ("23/07/2026", "7 tacos + 3 aguas + propina", "Pedro", 210, "efectivo", "propina $30"),
    ("19/07/2026", "5 tacos al pastor grandes", "Juan", 140, "efectivo", ""),
    ("22/07/2026", "3 tacos + coca", "Doña Rosa", 90, "efectivo", ""),
]

for row in ventas_data:
    ws.append(row)

# Agregar notas sueltas en celdas random (como haría alguien desordenado)
ws["H3"] = "← checar este"
ws["H8"] = "URGENTE: comprar servilletas"
ws["H12"] = "llamar proveedor tortillas"
ws["H17"] = "META: $3000/día"

# Formato "casero" — algunos headers en negritas, otros no
for col in range(1, 7):
    ws.cell(row=1, column=col).font = Font(bold=True, size=11)

# Columna de totales mal calculada abajo
ws.append([])
ws.append(["", "", "TOTAL (creo)", "=SUMA(D2:D21)", "", "no estoy segura"])

# ============================================================
# HOJA 2: "Clientes" — lista informal
# ============================================================
ws2 = wb.create_sheet("Clientes")
ws2.append(["Nombre", "Tel", "Que piden siempre", "Debe?"])
clientes = [
    ("Don Pedro", "55-1234-5678", "tacos pastor", "no"),
    ("Familia Lopez", "55-9876-5432", "orden familiar", "si $280"),
    ("Carlos", "", "tortas", "no"),
    ("Miguel", "55-4444-3333", "tacos + gringas", "no"),
    ("Ana", "55-1111-2222", "quesadillas", "no"),
    ("Sra Martinez", "55-3333-4444", "para llevar", "no"),
    ("Roberto", "55-5555-6666", "gringas", "no"),
    ("Laura", "55-7777-8888", "tortas + tacos", "checar"),
    ("Lic. Gomez", "55-9999-0000", "eventos", "si $150"),
    ("vecina Paty", "", "quesadillas", "si $5"),
    ("Doña Rosa", "55-2222-1111", "tacos chicos", "no"),
    ("Juan", "", "pastor grandes", "no"),
]
for c in clientes:
    ws2.append(c)

# ============================================================
# HOJA 3: "Gastos" — egresos mezclados
# ============================================================
ws3 = wb.create_sheet("Gastos")
ws3.append(["Fecha", "Qué compré", "Cuánto", "Dónde", "Nota"])
gastos = [
    ("18/07/2026", "Tortillas 5kg", 80, "Tortillería don Juan", ""),
    ("18/07/2026", "Carne pastor 3kg", 450, "Carnicería La Fe", "subió $20"),
    ("19/07/2026", "Refrescos (caja)", 180, "Abarrotera", ""),
    ("19/07/2026", "Aguas (paquete 24)", 95, "Costco", ""),
    ("20/07/2026", "Gas LP tanque", 550, "Gasero", "URGENTE necesito otro"),
    ("20/07/2026", "Servilletas + vasos", 120, "Papelería", ""),
    ("21/07/2026", "Tortillas 5kg", 80, "Tortillería don Juan", ""),
    ("21/07/2026", "Carne bistec 2kg", 380, "Carnicería La Fe", ""),
    ("21/07/2026", "Limones 2kg + cebolla", 65, "Mercado", ""),
    ("22/07/2026", "Tortillas 5kg", 80, "Tortillería don Juan", ""),
    ("22/07/2026", "Carne pastor 4kg", 600, "Carnicería La Fe", ""),
    ("22/07/2026", "Salsa + verduras", 150, "Mercado", ""),
    ("23/07/2026", "Tortillas 5kg", 80, "Tortillería don Juan", ""),
    ("23/07/2026", "Agua embotellada", 110, "Abarrotera", ""),
    ("23/07/2026", "Horchata preparada 10L", 180, "Sra. Carmen", ""),
]
for g in gastos:
    ws3.append(g)

ws3.append([])
ws3.append(["", "Total semana (aprox)", "=SUMA(C2:C16)", "", "no incluye propinas"])

# ============================================================
# HOJA 4: "Pendientes" — notas sueltas
# ============================================================
ws4 = wb.create_sheet("Pendientes!!")
ws4.append(["PENDIENTES DE LUPITA"])
ws4.append([])
ws4.append(["- Cobrar a Lic. Gomez ($150)"])
ws4.append(["- Cobrar a Familia Lopez ($280)"])
ws4.append(["- Dar cambio a vecina Paty ($5)"])
ws4.append(["- Comprar más gas (se acaba el viernes)"])
ws4.append(["- Pedir más carne, subió de precio"])
ws4.append(["- Checar si Laura pagó bien"])
ws4.append(["- META: vender $3000 diarios"])
ws4.append(["- Preguntar a contador por factura de Lic Gomez"])
ws4.append([])
ws4.append(["¿CUÁNTO VENDÍ ESTA SEMANA??? 🤷‍♀️"])
ws4.append(["(no sé, tendría que sumar todo a mano...)"])

# Ajustar anchos de columna para que se vea "real"
for ws_obj in [ws, ws2, ws3, ws4]:
    for col in range(1, 9):
        ws_obj.column_dimensions[get_column_letter(col)].width = 18

# Guardar
output_path = r"d:\02-A\code\bootcamp\demo_taqueria_lupita.xlsx"
wb.save(output_path)
print(f"✅ Excel generado: {output_path}")
print(f"   - Hoja 'Ventas julio': 20 ventas desordenadas")
print(f"   - Hoja 'Clientes': 12 clientes informales")
print(f"   - Hoja 'Gastos': 15 egresos")
print(f"   - Hoja 'Pendientes!!': notas caóticas")
